from pathlib import Path
import re
from difflib import SequenceMatcher

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import generate_novel_framework as source
from generate_novel_framework_clean import LEVELS


TEMPLATE = Path("/Users/wanglibo/Downloads/番茄小说/小说模板.docx")
MANUAL = Path("/Users/wanglibo/Desktop/GPT/manual_chapter_addons.txt")
OUTPUT = Path("/Users/wanglibo/Desktop/GPT/我都无敌了你们才来认亲_人工去重终稿.docx")


FORBIDDEN = [
    "本章将", "本章目标", "关系推进", "核心爽点", "情绪价值", "读者预期",
    "主导线", "承载事件", "完成任务", "反审判", "阶段末", "关系主轴",
    "赢下拳头", "策划术语", "办案式", "写作提醒", "硬约束", "指标",
    "关系链", "共同定义", "故事节奏", "查案", "调查搭档", "核心线",
    "主线矛盾", "自然事件", "关系一变", "反派使用原则", "本章重点",
]


def clear_document(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    for name, size in (("Title", 20), ("Subtitle", 11), ("Heading 1", 16), ("Heading 2", 13)):
        try:
            style = doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)


def paragraph(doc, text="", style=None):
    return doc.add_paragraph(text, style=style)


def labeled(doc, name, value):
    para = doc.add_paragraph()
    run = para.add_run(name)
    run.bold = True
    para.add_run(value)


def read_manual():
    result = {}
    for line in MANUAL.read_text(encoding="utf-8").splitlines():
        number, text = line.split("|", 1)
        result[int(number)] = text
    if set(result) != set(range(1, 201)):
        raise ValueError("人工章节内容必须正好包含第1至第200章")
    return result


def sentences(text):
    return [part for part in re.split(r"(?<=[。！？])", text) if part.strip()]


def similarity(left, right):
    # Chinese sentences没有空格，用字符序列比较重复程度比按词拆分更稳定。
    return SequenceMatcher(None, left, right).ratio()


def merge_chapter(event_body, addon, ending):
    main = sentences(event_body)
    selected = list(main)
    for sentence in sentences(addon):
        if max((similarity(sentence, old) for old in selected), default=0) < 0.48:
            selected.append(sentence)
    if max((similarity(ending, old) for old in selected), default=0) < 0.58:
        selected.append(ending)
    return "".join(selected)


# 这些章节原有两段素材的动作重复较多，改成一段连贯的作者叙述。
MANUAL_REWRITES = {
    1: "母亲葬礼后的雨夜，林烬在旧宅灵堂前醒来，玄烬天帝的记忆连同无边力量一起归位。他把气息压在引灵境，带着林小满走到门口。林岳拿族规逼兄妹搬走，抬手便要抢走母亲的木盒，林烬只把茶盏放到他面前，淡淡道：“钥匙留下，人可以走。”几名护卫膝盖砸进青砖，林岳攥紧的族牌也当场裂开。苏晚撑伞站在雨里，看着林小满抱紧木盒，第一次没有劝他们忍下这口气。墙里传来沉闷锁响，像有什么东西在旧宅深处醒了。",
    2: "青禾高中复测台前，班主任把林烬从武道班划进杂务名单，顾青河又故意撞翻苏晚的药箱，逼他当众低头。林烬伸指按住检测石，碎掉的灵光重新聚拢，顾青河的拳头停在他眉前，膝盖却先重重砸地。班主任赶紧收回杂务牌，围观学生盯着重新亮起的检测石，谁也没想到那个被嘲笑的废灵者会让武院检测员亲自下车行礼。顾青河脸上的笑一点点消失，校门外的黑车却已经点名要见林烬。",
    7: "林岳在夜市摆下赌台，把林烬母亲留下的旧钥匙当成彩头，叫来几个年轻武者看笑话。林烬坐在最末席，任他们赢了三局，直到那只手伸向钥匙，他才屈指一弹。赌台从中间裂开，筹码整整齐齐落回原处，林岳的帮手跪在碎木板旁，把昧下的钱一枚枚推还给摊主。旧钥匙被林烬隔空接住，蓝光映出一座陌生高楼，林岳想抢，手还没碰到钥匙便被压趴在桌边。夜市重新响起人声，林烬掌中的蓝光却迟迟没有熄灭。",
    12: "林烬到云岚武院报到，教务处仍把他塞进器材库。顾青河让他背整车护具跑圈，自己抱臂站在旁边嘲笑。林烬单手托起器材车，转身把车轮放到顾青河面前，地面被余力震出一圈裂纹。顾青河刚想骂人，膝盖已经被无形力量压进尘土，教官只得把杂务牌换成正式训练牌。学生们不再围着顾青河起哄，反而有人悄悄递来拜帖。顾青河盯着新牌子咬牙说擂台上见，顾家的车也在校门外停了下来。",
    13: "顾青河在训练场摆下擂台，放话林烬只能接住自己三拳。第一拳擦着衣角过去，第二拳被林烬两指夹住，第三拳还没落下，顾青河已经被压到擂台边。裁判看着他发白的脸，连忙宣布林烬获胜。看台上原本等着看笑话的学生立刻围过来，问他能不能一起训练，顾青河却只能捡起掉在地上的挑战牌。顾家管家在场外等了许久，终于走到苏晚面前赔礼，再请林烬上车。车厢里的礼盒没有打开，黑潮气息却已经透了出来。",
    14: "老祝把退役城卫留下的短枪交给林烬，说林家最近一直在找一份旧档案。两个林家护卫闯进修理铺，伸手便要拿枪，林烬只在柜台上一拍，两人连同门板一起撞到巷外。老祝拆开枪托，里面露出半张城卫地图，终于承认当年见过林烬母亲在武院城墙下救人。林烬把地图收好，告诉护卫回去转告林岳，旧街的东西没有人能白拿。巷口那辆盯梢的黑车随即掉头，像是终于知道这里不该再来。",
    15: "林岳带着真罡武者夜闯旧宅，声称林烬偷走了林家的气运。威压落下，院中石桌先一步炸裂，林小满攥紧衣角。林烬从屋檐下走出来，抬手一按，来人的膝盖同时陷进青砖，碎掉的雨水在他脚边分成两道。真罡武者连兵器都拿不稳，只能挡在林岳身前求饶。林烬让他把林岳带走，留下话：“下一次再进门，就把鞋留下。”旧钥匙在林岳离开后自行转向武院，木盒上的母亲姓氏也跟着亮了一下。",
    16: "武院要求所有杂务生参加力量测试，顾青河故意把测试桩调到真罡级，等着看林烬出丑。林烬只用引灵境的力道挥拳，测试桩安然无恙，后方防护墙却一层层亮起。记录员许放盯着数值发呆，顾青河扑过去想抢走测试纸，林烬的鞋尖落在他手背旁，地砖立刻压出细纹。顾青河不敢再动，许放把记录折进袖中，低声告诉林烬，武院有人正在删改旧档案。周围的笑声从此停了，教官也换了称呼。",
    17: "林家准备拆掉旧街，韩家武者挨家挨户逼居民在搬迁契上按手印。林烬来到签字墙前，领头人把老祝推到一边，笑说一个引灵境管不了整条街。林烬抬掌拍在墙上，契约背后的隐藏条款一行行浮出，领头人还没看完便被按跪在桌旁。居民撕掉合约，老祝带头把赔偿数目写满整面墙。巷口的人终于不再躲着看，连原本准备帮韩家搬东西的武者也悄悄放下了手里的箱子。",
    19: "林烬带苏晚来到武院禁区，顾青河带着天才班堵在门前，拿院规说普通学员不得靠近。苏晚被人推得后退半步，林烬一拳砸在禁区门的封印石上，锁链齐齐崩断，挡路的导师也跪成一排。院长赶到时，顾青河正捡起掉在地上的封锁牌，林烬让他把牌子挂回原处，再亲手把路让开。黑雾从门缝里涌出来，苏晚听见父亲的声音，却只听清半句：“不要让他坐上王座。”更深处随即有人喊出玄烬的名字。",
    20: "顾家长辈带着礼物到旧宅赔礼，嘴上说顾青河只是年轻气盛，手却一直按着袖中短刃。林烬让他把短刃放到桌上，一指压下，刃口立刻弯成圆环。顾家人收起笑脸，补齐武院欠下的资源，又让顾青河跪下向苏晚道歉。苏晚不要礼物，只要被扣的学生名额，林小满也把自己的要求说得清清楚楚。顾家车队刚走，城卫署的黑潮警报便响遍全城，旧钥匙在林烬掌心发热，雨幕外的街道同时亮起红灯。",
}


def chapter_text(stage, chapter_no, event, manual):
    title, body, ending = event
    story = MANUAL_REWRITES.get(chapter_no, merge_chapter(body, manual[chapter_no], ending))
    return (
        f"第{chapter_no}章《{stage['name']}·{title}》\n"
        f"{story} 林烬当前对外显露的实力，是{stage['public']}。"
    )


def build():
    manual = read_manual()
    doc = Document(str(TEMPLATE))
    clear_document(doc)
    setup_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = paragraph(doc, "小说模板：我都无敌了你们才来认亲", "Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(doc, "百万字现代都市高武日常男频无敌打脸爽文框架", "Subtitle")

    paragraph(doc, "项目信息：", "Heading 1")
    labeled(doc, "书名：", "《我都无敌了你们才来认亲》")
    labeled(doc, "类型：", "现代都市、高武、日常、无敌、打脸、男频爽文")
    labeled(doc, "风格：", "语言自然利落，场景从早餐店、旧街、校门和训练场开始。林烬很少解释，别人伸手抢东西，他就当场把东西拿回来；别人挡住妹妹和街坊的路，他就走过去让对方让开。每次交手都留下看得见的结果：有人跪下赔礼，有人把位置让出来，也有人递上战帖，下一场冲突紧跟着发生。")
    labeled(doc, "预计字数：", "约100万字；10个阶段，每阶段约10万字；共200章，每章正文约5000字。")
    labeled(doc, "故事基调：", "林烬在第一章雨夜觉醒前世玄烬天帝的记忆和真正力量，平日只让外界看见不同的境界。他带着妹妹从旧宅、学校和旧街走出去，靠一次次当面出手把房子、学籍、座位、粮食和尊严拿回来。林烬不会受伤，也不会真正落败。")

    paragraph(doc, "故事核心：", "Heading 1")
    paragraph(doc, "母亲葬礼后的雨夜，林烬醒来，记起自己曾是玄烬天帝。他带着妹妹住进被林家赶出的旧宅，顾青河扔来杂务牌，周铁带人收街费，林岳拿族规抢房契，林烬便在校门、旧街和家宴上一一出手。后来黑潮封城，王庭、圣城、帝国和星海相继登场，曾经让他站在门外的人，最后都要亲自把门打开。")

    paragraph(doc, "一、圣经：", "Heading 1")
    labeled(doc, "小说圣经：", "这是一个武力决定位置的现代高武世界。林烬第一章就拥有无界天极的真实力量，之后只改变外界看见的境界，不削弱自身实力。有人拿走林小满的学籍，他让对方当场送回来；有人踢翻旧街摊子，他让对方跪在雨里赔钱；有人拿王位和家世压人，他就站到擂台中央，让对方自己把手里的东西放下。每场交手都落到具体的人和地方，冲突过后，座位、称呼、规矩或手里的东西一定会变。")
    labeled(doc, "世界观/规则：", "灵气复苏后，城市里有武院、武馆、城卫和世家，黑潮裂口会带来异兽，也会把埋藏多年的旧事推到人前。武院按境界分训练场，城卫按战功分守城令，世家和王庭把强弱写在座位、门牌和兵器上。强者可以抢资源，但只有真正站到普通人身前、让对手当众收手的人，才会得到长久的尊重。")
    labeled(doc, "章节规划：", "前二十章写旧宅、学校和旧街；二十一至四十章进入云岚武院；四十一至六十章守住黑潮中的云岚市；六十一至八十章打进九家会馆和林氏祖地；八十一至一百章走上北境国门；一百零一至一百二十章进入九州王庭；一百二十一至一百四十章进入圣城学宫；一百四十一至一百六十章进入大荒帝国；一百六十一至一百八十章驶入星海；一百八十一至二百章站在太初门前。第200章林烬走进陌生世界，门外的人各自守住未完的事，门内的新战帖刚刚落下。")

    paragraph(doc, "高武等级划分：", "Heading 1")
    for name, desc in LEVELS:
        labeled(doc, f"{name}：", desc)
    labeled(doc, "林烬的真实实力：", "无界天极。第一章觉醒即拥有这一层力量，之后从未削弱，也没有真正败过。")

    paragraph(doc, "主要人物：", "Heading 1")
    labeled(doc, "男主：", "林烬。前世是玄烬天帝，今生是被林家赶出旧宅的青年。话少，动手干净，护着妹妹和旧街的人。")
    labeled(doc, "家人与同伴：", "林小满，林烬的妹妹，体内藏着太初血印，从抱着木盒躲在哥哥身后，成长为亲手决定血印去留的人；苏晚，校医实习生，带着父亲留下的药箱靠近林烬，后来成为裂口研究者；楚明月，圣城圣女，从试探林烬到与他并肩守住诸界议席；韩月，九家会馆的掌印者，从旧婚约中走出来，替旧街和北境送出资源。")
    labeled(doc, "重要人物：", "顾青河、林岳、周铁、陆沉舟、叶惊鸿、宁王、霍真、白秋生、温策、萧无尘、顾长夜、阿九、洛夷、铁七、无咎等人，都会在出手和让位之间改变对林烬的态度。")

    for stage in source.STAGES:
        no = stage["no"]
        doc.add_page_break()
        paragraph(doc, f"第{no}阶段：{stage['name']}", "Heading 1")
        labeled(doc, "章节范围：", f"{stage['range']}；每章正文约5000字。")
        labeled(doc, "主要地点：", stage["place"])
        labeled(doc, "境界划分：", f"这一段故事里，林烬当前对外显露的实力为{stage['public']}。")
        labeled(doc, "阶段概要：", stage["summary"])
        labeled(doc, "人物变化：", stage["relations"])
        paragraph(doc, f"第{no}阶段人物", "Heading 2")
        for name, identity, goal, secret, voice, note in stage["characters"]:
            labeled(doc, "姓名：", name)
            labeled(doc, "身份：", identity)
            labeled(doc, "人物目标：", goal)
            labeled(doc, "隐藏秘密：", secret)
            labeled(doc, "语言风格：", voice)
            labeled(doc, "备注：", note)
            paragraph(doc)
        paragraph(doc, f"第{no}阶段章节细纲", "Heading 2")
        for offset, event in enumerate(source.STAGE_EVENT_DATA[no], 1):
            number = (no - 1) * 20 + offset
            paragraph(doc, chapter_text(stage, number, event, manual))

    doc.add_page_break()
    paragraph(doc, "二、设定：", "Heading 1")
    settings = [
        ("旧钥匙", "物品", "林烬母亲留下的线索被分成数段，藏在药箱、护腕、军令和星图里。它会把林烬带到下一处擂台，也会让更多人不得不说出当年藏起来的话。"),
        ("黑潮裂口", "事件", "连接不同区域和世界的裂口，会放出异兽，也会把被遮住的武者、军令和旧事推到众人面前。"),
        ("云岚武院", "地点、势力", "现代高武学校。林烬从挂着杂务牌进入，最后让院长亲自递上学员令。训练场、食堂、校门和禁区都可能成为擂台。"),
        ("太初血印", "能力、物品", "藏在林小满体内的印记，可以打开或关闭太初门。林小满会自己决定它的去留，不会只等着别人保护。"),
        ("太初门", "地点、规则", "门后连接未知世界。林烬走进去以后，门外的人仍要自己守住航道、议席和血印，门内的新敌人也不会因为旧事结束而消失。"),
        ("开放式结尾", "结局", "第200章林烬进入陌生世界，城墙上的武者把长枪对准他，门外通讯石仍然亮着。苏晚、林小满、楚明月和留下的人各有要守的地方，新的战帖落下，故事没有替任何人把以后的人生写死。"),
    ]
    for name, kind, desc in settings:
        labeled(doc, "名称：", name)
        labeled(doc, "类型：", kind)
        labeled(doc, "说明：", desc)
        paragraph(doc)

    full_text = "\n".join(p.text for p in doc.paragraphs)
    chapters = [p.text for p in doc.paragraphs if p.text.startswith("第") and "章《" in p.text]
    if len(chapters) != 200:
        raise ValueError(f"章节数量错误：{len(chapters)}")
    if len(set(chapters)) != 200:
        raise ValueError("章节内容存在整段重复")
    found = [word for word in FORBIDDEN if word in full_text]
    if found:
        raise ValueError(f"发现禁用措辞：{found}")
    lengths = [len(item) for item in chapters]
    if min(lengths) < 190 or max(lengths) > 360:
        raise ValueError(f"章节字数异常：{min(lengths)}-{max(lengths)}")
    for stage in source.STAGES:
        expected = f"林烬当前对外显露的实力，是{stage['public']}。"
        selected = chapters[(stage["no"] - 1) * 20: stage["no"] * 20]
        if any(expected not in item for item in selected):
            raise ValueError(f"第{stage['no']}阶段境界标记缺失")
    if "林烬当前对外显露的实力，是太初境。" not in chapters[-1]:
        raise ValueError("第200章未保留太初境标记")

    doc.core_properties.title = "我都无敌了你们才来认亲_人工去重终稿"
    doc.core_properties.subject = "现代都市高武日常男频无敌打脸小说框架"
    doc.core_properties.author = ""
    doc.save(str(OUTPUT))
    print(f"created: {OUTPUT}")
    print(f"chapters: {len(chapters)}")
    print(f"outline lengths: {min(lengths)}-{max(lengths)}")
    print("forbidden wording: none")


if __name__ == "__main__":
    build()
