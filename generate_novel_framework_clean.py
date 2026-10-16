from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import generate_novel_framework as legacy

from generate_novel_framework import STAGES, STAGE_EVENT_DATA


TEMPLATE = Path("/Users/wanglibo/Downloads/番茄小说/小说模板.docx")
OUTPUT = Path("/Users/wanglibo/Desktop/GPT/我都无敌了你们才来认亲_重新生成版.docx")


STAGE_CAST = {
    1: [
        ("林小满", "抱紧母亲的木盒，不再躲到林烬身后"),
        ("苏晚", "收好药箱，替旧街的人把门口的路让出来"),
        ("老祝", "把修理铺的旧门锁换掉，站到街坊前面"),
        ("周铁", "跪在巷口退钱，吐出林家藏着的消息"),
        ("顾青河", "咬着牙收起挑战牌，第一次不敢抢先开口"),
    ],
    2: [
        ("白棠", "把新名单贴上公告墙，叫人给林烬打开训练场"),
        ("顾宁", "关上顾家的车门，把被改掉的学籍重新递回去"),
        ("罗教头", "提着钥匙站在场边，亲口叫林烬的名字"),
        ("许放", "把监控备份藏进记录板，不再替天才班遮掩"),
        ("顾青河", "捡起落地的挑战书，低声问林烬何时再打"),
    ],
    3: [
        ("陆沉舟", "把通讯器交给林烬，带队守住最危险的城墙"),
        ("阮青", "挂起居民名单，带着旧街人走进防线"),
        ("沈砚", "转开镜头，把完整的救援画面播给全城"),
        ("白鹭", "拉开实验室的门，让林小满看见血印的变化"),
        ("马魁", "跪到配给台前，把扣下的药和粮食一袋袋推回来"),
    ],
    4: [
        ("韩月", "把家主印按在桌上，替林小满挡开追来的护卫"),
        ("林太岳", "走下祖地高台，把重审母亲旧案的令牌交出来"),
        ("唐七", "收起竞价牌，亲自把会馆的暗门打开"),
        ("林岳", "站在祖地门外等候，不敢再拿族规压人"),
        ("陆问川", "把九家贵宾令送到车前，连称呼也换了"),
    ],
    5: [
        ("叶惊鸿", "把守城令压在林烬手里，亲自带老兵退下"),
        ("谢临风", "收起自己的长剑，替林烬挡住境外代表"),
        ("秦昭", "收回招揽令，先让武盟把城门外的伤员送回去"),
        ("伊莎", "撕掉天门坐标的价格牌，叫海外武者放下兵器"),
        ("陈渡", "把原始战绩写回公示墙，让所有人看清每一拳"),
    ],
    6: [
        ("霍真", "把守军旗插到林烬身边，带兵打开王庭宫门"),
        ("宋离", "抱着旧档案，把被删掉的名字贴回宫墙"),
        ("宁王", "收起摄政冠，坐到守军和百姓都能看见的位置"),
        ("林小满", "握住发热的血印，自己走过王座后的暗门"),
        ("贺连城", "从高台退到宫门边，再也不敢煽动人群"),
    ],
    7: [
        ("楚明月", "收回圣女令，站在白玉长街上替普通学生说话"),
        ("白秋生", "收起长剑，替榜尾学生守住试炼台入口"),
        ("乔北", "抱着新榜单跑过长街，把末席名字一一改回去"),
        ("温策", "交出副院长印，亲手打开藏经阁侧门"),
        ("虞红衣", "把商会的资源送进学宫，不再向圣族低头"),
    ],
    8: [
        ("顾长夜", "交出浮空战城的权限，站到军团议席后面"),
        ("叶知秋", "带着普通人的席位走进帝国大殿"),
        ("沈小刀", "抱着扳手钻进核心舱，骂着修好失控的战城"),
        ("萧无尘", "收起储君印，把护卫交给真正守城的人"),
        ("阿九", "关掉议庭的炮口，带林烬走进战城下层"),
    ],
    9: [
        ("洛夷", "打开主港大门，把地球代表请进圆桌"),
        ("铁七", "撕掉悬赏单，带着护航舰队挡住追兵"),
        ("陆小满", "抱着通行册分发座位，让普通人也听懂星海话"),
        ("纪星河", "当众说出旧错，重新带队走完地球航道"),
        ("无咎", "踏上无名旧舰，把没有星图的黑暗指给林烬"),
    ],
    10: [
        ("苏晚", "握住通讯石，把门外的声音送到林烬耳边"),
        ("林小满", "把血印留在门框上，自己选留下还是跟随"),
        ("楚明月", "把诸界的兵器摆到门前，逐个收下承诺"),
        ("霍真", "带守军护住旗阵，不让任何人替别人做决定"),
        ("无咎", "站在门后报出陌生世界的战鼓声"),
    ],
}


STAGE_PROLOGUES = {
    1: "林烬在母亲葬礼后的雨夜醒来，玄烬天帝的记忆连着无边力量一起归位。他把气息压在引灵境，带着妹妹住进旧宅，亲眼看着林家、学校和旧街的人轮番伸手。",
    2: "林烬拿着正式学员令走进云岚武院，教务处却把杂务牌挂到他胸前。顾青河带着天才班一遍遍找上门，武院的每一扇门都在一次出手后换了方向。",
    3: "黑潮裂口撕开云岚市，城卫署把旧街挡在防线外。林烬带着街坊走上城墙，陆沉舟和马魁先后站到他面前，城门外的异兽替所有人问出了答案。",
    4: "云岚九家把母亲的遗物摆上会馆拍卖台，还拿族谱和婚约逼林烬回头。他从最末席一路打到祖地石门，韩月和林太岳各自看见了家族该向谁低头。",
    5: "交流队把林烬塞进北境补给车最后一排，境外天才隔着擂台笑他没有战绩。雪原上的兽潮撞响城门后，战旗和守城令一起落到他手中。",
    6: "宁王把王位和封号摆到林烬面前，贵族子弟却在王城拿平民守军取乐。林烬从演武场打进旧档案库，让霍真和宋离把被藏起来的名字带回大殿。",
    7: "圣城学宫先查血统，再看天骄榜，只给林烬一张最末席的木牌。十座试炼台前，白秋生守着长剑，楚明月守着自己的圣女令，普通学生也开始抬头看擂台。",
    8: "大荒帝国的皇子、军团和黑潮议庭同时向林烬伸手，浮空战城的炮口也对准了皇城。他一路从街头打到核心舱，所有想拿普通人垫脚的人都得先接住他的拳头。",
    9: "九曜舰环把地球代表赶进货舱，贵族巡卫还要收走航道和武器。林烬带着众人从货梯走进主港，在星海擂台和舰队甲板上替地球拿回座位。",
    10: "太初门外插满诸界王旗，所有人都想让林烬替自己坐上王座。苏晚、林小满和楚明月站在门前，林烬却把决定留给每一个还没有放下兵器的人。",
}


STAGE_CLOSINGS = {
    1: "旧钥匙在林烬掌心发热，武院方向传来一声钟响。",
    2: "禁区深处传来呼唤玄烬的回声，院长的正式学员令也在此时送到。",
    3: "城墙暂时安静，九家会馆的红色信号却在城市尽头亮起。",
    4: "北境战鼓穿过会馆，林小满掌心的血印第一次映出王城轮廓。",
    5: "天门石碑上的坐标指向九州王城，宁王的车队穿过雪原而来。",
    6: "王座后的暗门打开，圣城飞车停在王城上空，白秋生的挑战书随风落下。",
    7: "大荒帝国的战舰落在学宫门外，母亲留下的记录指向更远的星海。",
    8: "九曜舰环的图案从帝印裂纹里亮起，皇城之外的入港航道却突然关闭。",
    9: "无名旧界的钟声传遍星海，旧舰舷窗外浮出太初门和诸界王旗。",
    10: "林烬踏上陌生世界的石阶，城墙上的长枪齐齐抬起，新的战鼓在门后响起。",
}


LEVELS = [
    ("引灵境", "能感知灵气并引入体内，强化身体和基础武技。"),
    ("锻骨境", "骨骼经过灵气淬炼，力量、速度和承受力大幅提高。"),
    ("开脉境", "灵脉贯通，灵力可以离体形成攻击，也能护住周身。"),
    ("真罡境", "真罡外放，单人便能镇压低阶黑潮兽和小型裂口。"),
    ("天门境", "精神天门打开，感知和出手可以越过很远的距离。"),
    ("王境", "个人领域成形，能够调动一城灵力和大型战阵。"),
    ("圣境", "触及法则本源，一言一动都能改变局部规则。"),
    ("帝境", "帝道秩序展开，力量足以影响国家和星球级格局。"),
    ("神域境", "神域覆盖星海，可以与大型世界的力量正面相接。"),
    ("太初境", "触及世界起源，能够打开或关闭连接诸界的门。"),
    ("无界天极", "十境之上的真实战力，无法被常规测量，只属于林烬前世与今生。"),
]


def clear_document(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    for name, size in (("Title", 20), ("Subtitle", 11), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        try:
            style = doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)


def paragraph(doc, text="", style=None):
    return doc.add_paragraph(text, style=style)


def labeled(doc, label, value):
    p = doc.add_paragraph()
    run = p.add_run(label)
    run.bold = True
    p.add_run(value)


AFTERMATH_PATTERNS = [
    "{person}从人群里走出来，先把刚才落在地上的东西收好，又看了眼林烬离开的方向。{title}留下的动静还在场外传着，没人敢像刚才那样大声说话。",
    "{person}没有跟着旁人起哄，走到林烬面前低声说了几句，随后把被撞歪的桌椅重新摆正。门口的守卫看见这一幕，默默把兵器收回鞘中。",
    "人群散开一半，{person}还站在原地。{person}捡起那件被人踩过的物件，拍掉灰尘递回去，旁边原本想看热闹的人全都把目光移开了。",
    "{person}替受惊的人关好门，又回头看向林烬。刚才还堵在路上的人已经退到两边，脚下的空地足够让几个人并肩走过去。",
    "场边有人想把刚才的事轻轻带过，{person}直接把那张记录拍在桌上，叫对方把名字和赔礼一起写清楚。林烬没有插话，只等对方低头。",
    "{person}从台阶上下来时，手里还握着刚才抢回来的牌子。{person}把牌子交给失主，又指了指林烬走远的背影，提醒众人别再挡路。",
    "林烬收手之后，{person}先看了看倒在墙边的人，再看向身旁的同伴。没人敢再冲上去，连刚才喊得最响的那个人也主动让开了门。",
    "{person}把散乱的物资一件件放回原处，嘴里还在数少了多少。等数到最后一件，原本围在四周的人已经排成一列，等着向林烬赔礼。",
    "刚才被推到后面的人重新站回台前，{person}替他把衣领理好，才转身对拦路的人说：“还有谁要试？”四周立刻安静下来。",
    "{person}没有接对方递来的好处，只把那只手推回去，要求先把被砸坏的地方修好。林烬抬眼看了一下，修理的人立刻跑出了门。",
    "风从场外吹进来，{person}按住桌上快被吹走的纸，顺手把最重要的一页递给林烬。纸上的字刚露出来，旁边的人便开始互相使眼色。",
    "{person}站到林烬身后时，前面那些挡路的人已经不敢回头。有人想替自己辩解，{person}只把刚才看见的经过说了一遍，对方顿时没了声音。",
    "林烬没有解释来龙去脉，{person}却把现场发生的事讲给后来赶到的人听。听到最后，连原本准备帮忙的人也转身去找真正该赔礼的人。",
    "{person}在门口停了一下，把属于自己那份东西收进怀里。林烬往前走，{person}便跟上去，身后的门随即关上，里面再没有人敢追出来。",
    "场上只剩下兵器碰地的声音，{person}弯腰把那件兵器推到一旁，给林小满和苏晚让出位置。两人走过时，刚才动手的人全都低下了头。",
    "{person}替林烬挡住围过来的人，却没有抢着说话，只把手里的东西递给他。林烬收下之后，{person}才回头把剩下的人一个个赶回原位。",
    "有人还想拿身份压人，{person}当场报出刚才的经过，连同那人拿走的东西也一并说了出来。林烬往前一步，所有声音立刻停了。",
    "{person}带着身边的人退到墙边，等林烬走到门口后，才把那扇被撞坏的门重新扶起来。门外有人递来赔礼，{person}只让他放在地上。",
    "事情落下之后，{person}没有急着离开，而是走到林烬刚才站过的地方看了一眼。地上的痕迹还在，附近的人也终于知道谁才是该让路的人。",
    "{person}把刚才听到的话告诉苏晚和林小满，三人顺着林烬的方向追了几步。后面又有人喊住他们，手里拿着新的拜帖，不敢再直接上前。",
]


def scene_person(stage_no, chapter_no, event):
    title, body, ending = event
    stage = STAGES[stage_no - 1]
    candidates = [name for name, *_ in stage["characters"] if name != "林烬"]
    candidates += [name for name, _ in STAGE_CAST[stage_no] if name != "林烬"]
    for name in candidates:
        if name in body or name in ending:
            return name
    return STAGE_CAST[stage_no][(chapter_no - 1) % len(STAGE_CAST[stage_no])][0]


def chapter_text(stage_no, chapter_no, event):
    title, event_body, ending = event
    stage = STAGES[stage_no - 1]
    person = scene_person(stage_no, chapter_no, event)
    aftermath = AFTERMATH_PATTERNS[(chapter_no - 1) % len(AFTERMATH_PATTERNS)].format(
        person=person,
        title=title,
    )
    return (
        f"第{chapter_no}章《{stage['name']}·{title}》\n"
        f"{event_body} {aftermath} {ending} "
        f"林烬对外显露的实力，是{stage['public']}。"
    )


def build_document():
    doc = Document(str(TEMPLATE))
    clear_document(doc)
    setup_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("小说模板：我都无敌了你们才来认亲")
    paragraph(doc, "百万字现代都市高武日常男频无敌打脸爽文框架", "Subtitle")

    paragraph(doc, "项目信息：", "Heading 1")
    labeled(doc, "书名：", "《我都无敌了你们才来认亲》")
    labeled(doc, "类型：", "现代都市、高武、日常、无敌、打脸、男频爽文")
    labeled(doc, "风格：", "句子利落，场景贴着人物走，出手、挨打、让路、赔礼都写清楚。林烬不站在旁边讲道理，谁挡住妹妹的路，他就走过去；谁伸手拿走街坊的东西，他就把东西原样放回去。日常里有早餐店、校门、训练场、旧街和城墙，打到更远的地方，也始终围着身边的人和一顿饭、一间屋、一条回家的路展开。")
    labeled(doc, "预计字数：", "约100万字；10个阶段，每阶段约10万字；每章约5000字，共200章。")
    labeled(doc, "故事基调：", "男主第一章就在雨夜想起前世和力量，之后始终没有真正输过，也没有人能伤到他。人人看不起他时，他带着妹妹在旧宅过日子；等他一次次出手把东西拿回来，亲戚、同学、武馆、世家和王庭便一次次换称呼、让位置、送东西。")

    paragraph(doc, "故事核心：", "Heading 1")
    paragraph(doc, "母亲葬礼后的雨夜，林烬醒来，记起自己曾是玄烬天帝。他不急着抬出真正的力量，只把气息压在最低处，带着林小满住进被林家赶出的旧宅。顾青河扔来杂务牌，周铁带人收街费，林岳拿族规抢房契，林烬便在校门、旧街和家宴上一一出手。后来黑潮封城，王庭、圣城、帝国和星海相继登场，曾经让他站在门外的人，最后都要亲自把门打开。")

    paragraph(doc, "一、圣经：", "Heading 1")
    labeled(doc, "小说圣经：", "林烬永远不受伤，不靠惨胜换来尊重。有人拿走林小满的学籍，他让对方当场送回来；有人踢翻旧街摊子，他让对方跪在雨里赔钱；有人拿王位和家世压人，他就站到擂台中央，让对方自己把手里的东西放下。每场交手都要有看得见的结果，或者让路，或者赔礼，或者把原先藏着的线索露出来。")
    labeled(doc, "世界观/规则：", "这是一个以武力决定位置的现代高武世界。武院按境界分训练场，武馆按拳头收街费，城卫按战功分守城令，世家、王庭和星海贵族也都把强弱写在座位、门牌和兵器上。黑潮裂口会带来异兽，也会让隐藏的力量和旧事浮出水面。境界决定别人敢不敢伸手，真正交手时，谁能站着把人护在身后，谁就有资格让所有人听他说话。")
    labeled(doc, "章节规划：", "前二十章写旧宅、学校和旧街，林烬带着妹妹把被夺走的房子、学籍和钱一件件拿回来；二十一章以后进入武院、封锁城、九家会馆和北境关城，认识他的人从冷眼看戏变成主动让路。王庭、圣城和大荒帝国把王位、席位和战城摆到面前，林烬只把这些位置推回守军和普通人手里。星海线揭开玄烬天帝未死的消息，太初门前的伙伴和旧敌各自留下，林烬走进门后的陌生世界，新的战鼓在结尾响起。")

    paragraph(doc, "高武等级与男主实力：", "Heading 1")
    for name, desc in LEVELS:
        labeled(doc, f"{name}：", desc)
    labeled(doc, "林烬：", "第一章觉醒时就拥有无界天极的真实力量，平日只让外人看见引灵境、锻骨境到太初境的不同气息。每一段故事里，他对外显露的境界不同，但真实实力没有被削弱，出手时不会受伤，也不会真正落败。")

    paragraph(doc, "主要人物：", "Heading 1")
    labeled(doc, "核心人物：", "林烬、苏晚、林小满、楚明月、顾宁。")
    labeled(doc, "人物关系：", "林家从把林烬赶出祖宅，到站在祖地门外等他回头；苏晚从替他收起药箱，到拿着旧钥匙追问他不肯说的过去；林小满从抱着木盒躲在他身后，到亲手把血印留在太初门上；楚明月从圣城长街的试探走到门前，和他一起看着诸界放下兵器。")
    labeled(doc, "对手与势力：", "顾青河、林岳、马魁、宁王、温策和黑潮议庭都会先用身份、境界或规矩压人，再在擂台、街口、城墙和王庭门前被林烬逼得低头。有的人低头后真正改过，有的人赔完礼还想再伸手，新的冲突会从他们的下一次出手开始。")

    for stage in STAGES:
        stage_no = stage["no"]
        doc.add_page_break()
        paragraph(doc, f"第{stage_no}阶段：{stage['name']}", "Heading 1")
        labeled(doc, "字数与章节：", f"{stage['range']}；每章约5000字。")
        labeled(doc, "地点：", stage["place"])
        labeled(doc, "对外显露的实力：", stage["public"])
        labeled(doc, "这一段故事：", STAGE_PROLOGUES[stage_no])
        labeled(doc, "人物变化：", stage["relations"])
        paragraph(doc, f"第{stage_no}阶段人物", "Heading 2")
        for name, identity, goal, secret, voice, note in stage["characters"]:
            labeled(doc, "姓名：", name)
            labeled(doc, "身份：", identity)
            labeled(doc, "人物目标：", goal)
            labeled(doc, "隐藏秘密：", secret)
            labeled(doc, "语言风格：", voice)
            labeled(doc, "备注：", note)
            paragraph(doc)

        paragraph(doc, f"第{stage_no}阶段章节细纲", "Heading 2")
        for offset, event in enumerate(STAGE_EVENT_DATA[stage_no], start=1):
            chapter_no = (stage_no - 1) * 20 + offset
            paragraph(doc, chapter_text(stage_no, chapter_no, event))

    doc.add_page_break()
    paragraph(doc, "二、设定：", "Heading 1")
    settings = [
        ("旧钥匙", "物品", "林烬母亲留下的线索被分成数段，藏在药箱、护腕、军令和星图里。每次有人为了它出手，林烬都会顺着打斗后的痕迹走向下一处地方。"),
        ("黑潮裂口", "事件", "连接不同区域和世界的裂口，会放出异兽，也会把被藏起来的人和旧事推到众人面前。"),
        ("云岚武院", "地点、势力", "现代高武学校，林烬从挂着杂务牌进入，到让院长亲自递上学员令。训练场、食堂、校门和禁区都能发生交手。"),
        ("九州王庭", "势力、地点", "掌握王城、军队和旧时代档案。宁王把王位摆在林烬面前，林烬却让守军和百姓走进王庭。"),
        ("太初血印", "能力、物品", "藏在林小满体内的印记，可以打开或关闭太初门。她会自己决定如何使用它，不会被当成只等保护的物件。"),
        ("太初门", "地点、规则", "门后连接未知世界，打开以后没有圆满答案，只有新的城墙、长枪和战鼓。林烬走进去，门外的人仍要自己守住自己的路。"),
        ("开放式结尾", "结局", "第200章林烬进入陌生世界，苏晚、林小满、楚明月和留下的人各自守住门外的事情。门内的新敌人刚举起长枪，门外的诸界也还在争论下一步怎么走。"),
    ]
    for name, kind, desc in settings:
        labeled(doc, "名称：", name)
        labeled(doc, "类型：", kind)
        labeled(doc, "说明：", desc)
        paragraph(doc)

    doc.core_properties.title = "我都无敌了你们才来认亲_重新生成版"
    doc.core_properties.subject = "现代都市高武日常男频无敌打脸小说框架"
    doc.core_properties.author = ""
    doc.save(str(OUTPUT))


def validate():
    doc = Document(str(OUTPUT))
    chapters = [p.text for p in doc.paragraphs if p.text.startswith("第") and "章《" in p.text]
    if len(chapters) != 200:
        raise AssertionError(f"chapter count: {len(chapters)}")
    lengths = [len(x) for x in chapters]
    if min(lengths) < 260 or max(lengths) > 390:
        raise AssertionError(f"chapter lengths: {min(lengths)}-{max(lengths)}")
    full_text = "\n".join(p.text for p in doc.paragraphs)
    forbidden = [
        "本章将", "本章目标", "关系推进", "核心爽点", "情绪价值", "读者预期",
        "主导线", "承载事件", "完成任务", "反审判", "阶段末", "关系主轴",
        "赢下拳头", "策划术语", "办案式", "写作提醒", "硬约束", "指标",
        "关系链", "共同定义", "故事节奏", "查案", "调查搭档", "核心线",
        "主线矛盾", "自然事件", "关系一变", "反派使用原则", "本章重点",
    ]
    found = [word for word in forbidden if word in full_text]
    if found:
        raise AssertionError(f"forbidden wording: {found}")
    if len(set(chapters)) != 200:
        raise AssertionError("duplicate chapter outline")
    for stage in STAGES:
        expected = f"林烬对外显露的实力，是{stage['public']}。"
        stage_chapters = chapters[(stage['no'] - 1) * 20:stage['no'] * 20]
        if any(expected not in item for item in stage_chapters):
            raise AssertionError(f"public realm missing in stage {stage['no']}")
    print(f"created: {OUTPUT}")
    print(f"chapters: {len(chapters)}")
    print(f"outline lengths: {min(lengths)}-{max(lengths)}")
    print("forbidden wording: none")


if __name__ == "__main__":
    build_document()
    validate()
