from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import generate_novel_framework as source
from generate_novel_framework_clean import LEVELS


TEMPLATE = Path("/Users/wanglibo/Downloads/番茄小说/小说模板.docx")
MANUAL = Path("/Users/wanglibo/Desktop/GPT/manual_chapter_addons.txt")
OUTPUT = Path("/Users/wanglibo/Desktop/GPT/我都无敌了你们才来认亲_全新人工生成版.docx")


FORBIDDEN = [
    "本章将", "本章目标", "关系推进", "核心爽点", "情绪价值", "读者预期",
    "主导线", "承载事件", "完成任务", "反审判", "阶段末", "关系主轴",
    "赢下拳头", "策划术语", "办案式", "写作提醒", "硬约束", "指标",
    "关系链", "共同定义", "故事节奏", "查案", "调查搭档", "核心线",
    "主线矛盾", "自然事件", "关系一变", "反派使用原则", "本章重点",
]


def clear_document(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    for name, size in (("Title", 20), ("Subtitle", 11), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        try:
            style = doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)


def p(doc, text="", style=None):
    return doc.add_paragraph(text, style=style)


def label(doc, name, text):
    para = doc.add_paragraph()
    run = para.add_run(name)
    run.bold = True
    para.add_run(text)


def read_manual():
    result = {}
    for line in MANUAL.read_text(encoding="utf-8").splitlines():
        number, text = line.split("|", 1)
        result[int(number)] = text
    if set(result) != set(range(1, 201)):
        raise ValueError("人工章节内容必须正好包含第1至第200章")
    if len(set(result.values())) != 200:
        raise ValueError("人工章节内容存在整段重复")
    return result


def chapter_text(stage, chapter_no, event, manual):
    title, body, ending = event
    return (
        f"第{chapter_no}章《{stage['name']}·{title}》\n"
        f"{body} {manual[chapter_no]} {ending} "
        f"林烬当前对外显露的实力，是{stage['public']}。"
    )


def build():
    manual = read_manual()
    doc = Document(str(TEMPLATE))
    clear_document(doc)
    styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = p(doc, "小说模板：我都无敌了你们才来认亲", "Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p(doc, "百万字现代都市高武日常男频无敌打脸爽文框架", "Subtitle")

    p(doc, "项目信息：", "Heading 1")
    label(doc, "书名：", "《我都无敌了你们才来认亲》")
    label(doc, "类型：", "现代都市、高武、日常、无敌、打脸、男频爽文")
    label(doc, "风格：", "语言自然利落，场景从早餐店、旧街、校门和训练场开始。林烬很少解释，别人伸手抢东西，他就当场把东西拿回来；别人挡住妹妹和街坊的路，他就走过去让对方让开。每次交手都留下看得见的结果：有人跪下赔礼，有人把位置让出来，也有人递上战帖，下一场冲突紧跟着发生。")
    label(doc, "预计字数：", "约100万字；10个阶段，每阶段约10万字；共200章，每章正文约5000字。")
    label(doc, "故事基调：", "林烬在第一章雨夜觉醒前世记忆和真正力量，真实战力始终凌驾诸境。他把气息压在外界看得懂的等级里，带着妹妹一路生活、出手、护人，从被赶出旧宅到让世家、王庭和星海主动让席。林烬不会受伤，也不会真正落败。")

    p(doc, "故事核心：", "Heading 1")
    p(doc, "母亲葬礼后的雨夜，林烬记起自己曾是玄烬天帝。他不急着亮出真正力量，只带着妹妹住进被林家赶出的旧宅。学校、武馆、城卫、世家和王庭轮番伸手，他便在校门、街口、擂台和城墙前一一出手，把被夺走的房子、学籍、钱、席位和尊严拿回来。后来黑潮裂口通向王庭、圣城、帝国和星海，曾经不肯让他进门的人，最后都要亲自打开门。")

    p(doc, "一、圣经：", "Heading 1")
    label(doc, "小说圣经：", "这是一个武力决定位置的现代高武世界。林烬第一章就拥有无界天极的真实力量，之后只改变外界看见的境界，不削弱自身实力。每章都有具体场景和交手，林烬不靠受伤、惨胜或突然失去力量制造悬念。")
    label(doc, "世界观与规则：", "灵气复苏后，城市里有武院、武馆、城卫和世家，黑潮裂口会带来异兽，也会把埋藏多年的旧事推到人前。境界高的人能抢到更好的训练场、座位和资源，但真正能守住普通人、让对手当众收回手的人，才会得到长久的尊重。")
    label(doc, "章节规划：", "前二十章写旧宅、学校和旧街；二十一至四十章进入云岚武院；四十一至六十章守住黑潮中的云岚市；六十一至八十章打进九家会馆和林氏祖地；八十一至一百章走上北境国门；一百零一至一百二十章进入九州王庭；一百二十一至一百四十章进入圣城学宫；一百四十一至一百六十章进入大荒帝国；一百六十一至一百八十章驶入星海；一百八十一至二百章站在太初门前。第200章林烬走进陌生世界，门外的人各自守住未完的事，门内的新战帖刚刚落下。")

    p(doc, "高武等级划分：", "Heading 1")
    for name, desc in LEVELS:
        label(doc, f"{name}：", desc)
    label(doc, "林烬的真实实力：", "无界天极。第一章觉醒即拥有这一层力量，之后从未削弱，也没有真正败过。")

    p(doc, "主要人物：", "Heading 1")
    label(doc, "男主：", "林烬。前世是玄烬天帝，今生是被林家赶出旧宅的青年。话少，动手干净，护着妹妹和旧街的人。")
    label(doc, "家人与同伴：", "林小满，林烬的妹妹，体内藏着太初血印，从抱着木盒躲在哥哥身后，成长为亲手决定血印去留的人；苏晚，校医实习生，带着父亲留下的药箱靠近林烬，后来成为裂口研究者；楚明月，圣城圣女，从试探林烬到与他并肩守住诸界议席；韩月，九家会馆的掌印者，从旧婚约中走出来，替旧街和北境送出资源。")
    label(doc, "重要人物：", "顾青河、林岳、周铁、陆沉舟、叶惊鸿、宁王、霍真、白秋生、温策、萧无尘、顾长夜、阿九、洛夷、铁七、无咎等人，都会在出手和让位之间改变对林烬的态度。")

    for stage in source.STAGES:
        no = stage["no"]
        doc.add_page_break()
        p(doc, f"第{no}阶段：{stage['name']}", "Heading 1")
        label(doc, "章节范围：", f"{stage['range']}；每章正文约5000字。")
        label(doc, "主要地点：", stage["place"])
        label(doc, "境界划分：", f"这一段故事里，林烬当前对外显露的实力为{stage['public']}。")
        label(doc, "阶段概要：", stage["summary"])
        label(doc, "人物变化：", stage["relations"])
        p(doc, f"第{no}阶段人物", "Heading 2")
        for name, identity, goal, secret, voice, note in stage["characters"]:
            label(doc, "姓名：", name)
            label(doc, "身份：", identity)
            label(doc, "人物目标：", goal)
            label(doc, "隐藏秘密：", secret)
            label(doc, "语言风格：", voice)
            label(doc, "备注：", note)
            p(doc)
        p(doc, f"第{no}阶段章节细纲", "Heading 2")
        for offset, event in enumerate(source.STAGE_EVENT_DATA[no], 1):
            number = (no - 1) * 20 + offset
            p(doc, chapter_text(stage, number, event, manual))

    doc.add_page_break()
    p(doc, "二、设定：", "Heading 1")
    settings = [
        ("旧钥匙", "物品", "林烬母亲留下的线索被分成数段，藏在药箱、护腕、军令和星图里。它会把林烬带到下一处擂台，也会让更多人不得不说出当年藏起来的话。"),
        ("黑潮裂口", "事件", "连接不同区域和世界的裂口，会放出异兽，也会把被遮住的武者、军令和旧事推到众人面前。"),
        ("云岚武院", "地点、势力", "现代高武学校。林烬从挂着杂务牌进入，最后让院长亲自递上学员令。训练场、食堂、校门和禁区都可能成为擂台。"),
        ("太初血印", "能力、物品", "藏在林小满体内的印记，可以打开或关闭太初门。林小满会自己决定它的去留，不会只等着别人保护。"),
        ("太初门", "地点、规则", "门后连接未知世界。林烬走进去以后，门外的人仍要自己守住航道、议席和血印，门内的新敌人也不会因为旧事结束而消失。"),
        ("开放式结尾", "结局", "第200章林烬进入陌生世界，城墙上的武者把长枪对准他，门外通讯石仍然亮着。苏晚、林小满、楚明月和留下的人各有要守的地方，新的战帖落下，故事没有替任何人把以后的人生写死。"),
    ]
    for name, kind, desc in settings:
        label(doc, "名称：", name)
        label(doc, "类型：", kind)
        label(doc, "说明：", desc)
        p(doc)

    full = "\n".join(x.text for x in doc.paragraphs)
    chapters = [x for x in doc.paragraphs if x.text.startswith("第") and "章《" in x.text]
    if len(chapters) != 200:
        raise ValueError(f"章节数量错误：{len(chapters)}")
    if len(set(x.text for x in chapters)) != 200:
        raise ValueError("章节内容存在重复")
    if any(word in full for word in FORBIDDEN):
        found = [word for word in FORBIDDEN if word in full]
        raise ValueError(f"发现禁用措辞：{found}")
    lengths = [len(x.text) for x in chapters]
    if min(lengths) < 210 or max(lengths) > 360:
        raise ValueError(f"章节字数异常：{min(lengths)}-{max(lengths)}")
    for stage in source.STAGES:
        expected = f"林烬当前对外显露的实力，是{stage['public']}。"
        selected = chapters[(stage["no"] - 1) * 20: stage["no"] * 20]
        if any(expected not in x.text for x in selected):
            raise ValueError(f"第{stage['no']}阶段境界标记缺失")

    doc.core_properties.title = "我都无敌了你们才来认亲_全新人工生成版"
    doc.core_properties.subject = "现代都市高武日常男频无敌打脸小说框架"
    doc.core_properties.author = ""
    doc.save(str(OUTPUT))
    print(f"created: {OUTPUT}")
    print(f"chapters: {len(chapters)}")
    print(f"outline lengths: {min(lengths)}-{max(lengths)}")
    print("forbidden wording: none")


if __name__ == "__main__":
    build()
